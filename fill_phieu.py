import docx
from docx.shared import Pt, RGBColor

def fill_worksheet():
    doc = docx.Document('Phieu_thuc_hanh_CTDL_GT_2026_10_09.docx')
    
    # Fill Table 0: Bang truy vet duong va thoi gian
    t0 = doc.tables[0]
    t0.rows[1].cells[1].text = '1 -> 2'
    t0.rows[1].cells[2].text = '1'
    t0.rows[1].cells[3].text = '2(1) + 1 = 3'

    t0.rows[2].cells[1].text = '1 -> 2 -> 4'
    t0.rows[2].cells[2].text = '2'
    t0.rows[2].cells[3].text = '2(2) + 1 = 5'

    t0.rows[3].cells[1].text = '1 -> 2 -> 4 -> 6'
    t0.rows[3].cells[2].text = '3'
    t0.rows[3].cells[3].text = '2(3) + 1 = 7'

    # Fill Table 1: Giao va hoan tac
    t1 = doc.tables[1]
    t1.rows[1].cells[1].text = '[1, 2]'
    t1.rows[1].cells[2].text = '[]'
    t1.rows[1].cells[3].text = '0'
    t1.rows[1].cells[4].text = '0'

    t1.rows[2].cells[1].text = '[2]'
    t1.rows[2].cells[2].text = '[1]'
    t1.rows[2].cells[3].text = '3'
    t1.rows[2].cells[4].text = '6'

    t1.rows[3].cells[1].text = '[1, 2]'
    t1.rows[3].cells[2].text = '[]'
    t1.rows[3].cells[3].text = '0'
    t1.rows[3].cells[4].text = '0'

    # Fill paragraphs
    doc.paragraphs[1].text = "Họ tên: Nguyễn Hoàng Nam          MSSV: N24DCDT051          Lớp: D24CQDT01-N"
    doc.paragraphs[2].text = "GitHub: https://github.com/sinhvien/Bo_Thuc_hanh_Julia_2026_10_09    Commit SHA: 587283c"

    # P5: 1. Mo ta luong ung dung
    doc.paragraphs[5].text = (
        "- Mục tiêu: Điều phối giao đơn độc lập từ kho 1 tới các điểm trong campus và quay về kho, "
        "tối đa hóa lợi ích trong ngân sách thời gian B (B = 8 với MINI001, B = 26 với N24DCDT051).\n"
        "- Đầu vào: don_hang.csv (id, diem_giao, thoi_gian_phuc_vu, loi_ich) và ban_do.csv (7 cạnh vô hướng trên 6 đỉnh).\n"
        "- Đầu ra: Danh mục đơn, kế hoạch tối ưu, chuỗi giao/hoàn tác (pending, completed, spent, earned) và tệp ket_qua.csv.\n"
        "- Sơ đồ truyền dữ liệu: CSV -> read_data -> Danh mục đơn & Bản đồ kề (adjacency) -> merge_sort (sắp xếp ID tăng) "
        "-> bfs_routes (tìm số cạnh ngắn nhất h từ kho 1) -> costs = Dict(id => 2h + s) -> knapsack (Quy hoạch động 0/1 chọn "
        "tập đơn tối ưu trong ngân sách B) -> priority_orders (sắp xếp giảm theo tỷ số b/cost) -> nạp hàng đợi vòng FIFO "
        "CircularQueue (pending) -> deliver! chuyển sang ngăn xếp Stack (completed) và undo! khôi phục ngược lại -> export_result (xuất ket_qua.csv)."
    )

    # P9, P10, P11: NV1 merge_sort
    doc.paragraphs[9].text = (
        "Ý tưởng/điều kiện áp dụng: Áp dụng chia để trị (Divide and Conquer). Nếu length(a) <= 1 trả về collect(a). "
        "Chia đôi tại mid = div(n, 2), đệ quy sắp 2 nửa left, right rồi trộn lại bằng 2 con trỏ i, j. Nếu before(right[j], left[i]) "
        "thì lấy right[j], ngược lại lấy left[i] để đảm bảo tính ổn định (stable). Áp dụng cho mọi dãy có định nghĩa quan hệ so sánh before."
    )
    doc.paragraphs[10].text = (
        "Lời gọi và kết quả kiểm thử: julia TH_2026_10_09_Kiem_thu.jl --phase=sort -> ĐẠT 8/8 phép kiểm tra. "
        "Ví dụ: merge_sort([3, 1, 2], (<)) trả về [1, 2, 3]; dãy gốc không đổi; hoạt động đúng trên dãy rỗng và 1 phần tử."
    )
    doc.paragraphs[11].text = (
        "Chi phí thời gian/bộ nhớ: Thời gian O(n log n) trong mọi trường hợp; Bộ nhớ phụ O(n) cho bộ đệm trộn và O(log n) ngăn xếp đệ quy."
    )

    # P13, P14, P15: NV2 binary_search
    doc.paragraphs[13].text = (
        "Ý tưởng/điều kiện áp dụng: Duy trì khoảng tìm kiếm [lo, hi]. Lấy mid = lo + div(hi - lo, 2). Nếu a[mid].id == id trả về a[mid]; "
        "nếu a[mid].id < id thu hẹp nửa phải lo = mid + 1; nếu a[mid].id > id thu hẹp nửa trái hi = mid - 1. Nếu lo > hi trả về nothing. "
        "Tiền điều kiện: Mảng a bắt buộc phải được sắp xếp tăng dần theo id, các mã id là duy nhất."
    )
    doc.paragraphs[14].text = (
        "Lời gọi và kết quả kiểm thử: julia TH_2026_10_09_Kiem_thu.jl --phase=search -> ĐẠT 5/5 phép kiểm tra. "
        "Tìm ID 2 trả về Order(2, 4, 1, 8); tìm ID 99 hoặc tìm trên mảng rỗng [] trả về nothing."
    )
    doc.paragraphs[15].text = (
        "Chi phí thời gian/bộ nhớ: Thời gian O(log n); Bộ nhớ phụ O(1) do sử dụng vòng lặp lặp lại, không đệ quy."
    )

    # P17, P18, P19: NV3 bfs_routes
    doc.paragraphs[17].text = (
        "Ý tưởng/điều kiện áp dụng: Duyệt theo chiều rộng xuất phát từ đỉnh nguồn (kho 1) sử dụng hàng đợi vòng FIFO CircularQueue. "
        "Khởi tạo distance toàn -1, parent toàn 0. Đánh dấu distance[v] = distance[u] + 1 và parent[v] = u ngay khi enqueue để mỗi đỉnh "
        "chỉ được thăm đúng 1 lần. Dựng routes bằng cách đi ngược parent từ đích về nguồn rồi reverse. Áp dụng cho đồ thị vô hướng không trọng số."
    )
    doc.paragraphs[18].text = (
        "Lời gọi và kết quả kiểm thử: julia TH_2026_10_09_Kiem_thu.jl --phase=bfs -> ĐẠT 7/7 phép kiểm tra. "
        "Khoảng cách từ 1 tới các đỉnh 1..6 là [0, 1, 1, 2, 2, 3]; chi phí 3 đơn mini là [3, 5, 7]."
    )
    doc.paragraphs[19].text = (
        "Chi phí thời gian/bộ nhớ: Thời gian O(V + E); Bộ nhớ O(V) cho hàng đợi, distance, parent; lưu trữ các tuyến đường tốn O(V^2)."
    )

    # P21, P22, P23: NV4 knapsack
    doc.paragraphs[21].text = (
        "Ý tưởng/điều kiện áp dụng: Quy hoạch động Ba lô 0/1. Bảng table[n+1, budget+1] lưu F(i, b) tại table[i+1, b+1]. "
        "Hàng i=0 và cột b=0 có giá trị 0. Với đơn i có chi phí c_i, lợi ích val: nếu c_i <= b thì "
        "table[i+1, b+1] = max(table[i, b+1], val + table[i, b - c_i + 1]), ngược lại bằng table[i, b+1]. Khi bằng nhau giữ phương án "
        "không lấy. Truy vết ngược từ n về 1 để xác định tập đơn được chọn. Áp dụng cho bài toán chọn đơn 0/1 với chi phí nguyên."
    )
    doc.paragraphs[22].text = (
        "Lời gọi và kết quả kiểm thử: julia TH_2026_10_09_Kiem_thu.jl --phase=dp -> ĐẠT 12/12 phép kiểm tra. "
        "Bộ mini ngân sách 8 cho tập [1, 2], chi phí 8, lợi ích tối ưu 14; kết quả DP khớp 100% với vét cạn."
    )
    doc.paragraphs[23].text = (
        "Chi phí thời gian/bộ nhớ: Thời gian O(n * B); Bộ nhớ O(n * B) cho bảng quy hoạch động; thời gian truy vết O(n)."
    )

    # P26: 4. Quyet dinh chon don
    doc.paragraphs[26].text = (
        "- Ngân sách: B = 8. Tập đơn được chọn: {1, 2}. Tổng thời gian: 3 + 5 = 8. Tổng lợi ích: 6 + 8 = 14.\n"
        "- Giải thích ô F(2, 8) tại table[3, 9]: Không lấy đơn 2 thì F(1, 8) = table[2, 9] = 6; Lấy đơn 2 (c_2 = 5, v_2 = 8) "
        "thì 8 + F(1, 8 - 5) = 8 + table[2, 4] = 8 + 6 = 14. Chọn max(6, 14) = 14. Nhánh lấy phải lấy từ hàng i-1 (table[2, :]) "
        "để đảm bảo đơn 2 chỉ được dùng tối đa 1 lần.\n"
        "- Bước truy vết từ F(3, 8): table[4, 9] == table[3, 9] (14 == 14) -> không chọn đơn 3; table[3, 9] > table[2, 9] (14 > 6) "
        "-> chọn đơn 2, b còn 8 - 5 = 3; table[2, 4] > table[1, 4] (6 > 0) -> chọn đơn 1, b còn 3 - 3 = 0. Tập chọn: [1, 2].\n"
        "- Đối chiếu vét cạn: 2^3 = 8 tập con; tập {1, 2} có chi phí 8 <= 8 và lợi ích 14 là nghiệm duy nhất tối ưu, trùng khớp 100% với DP."
    )

    # P29, P30, P31: 6. Dau ra va nhan xet
    doc.paragraphs[29].text = (
        "Đường dẫn CSV kết quả: data/MINI001/ket_qua.csv (và data/N24DCDT051/ket_qua.csv)        Link notebook: TH_2026_10_09_Pluto.jl"
    )
    doc.paragraphs[30].text = (
        "Lỗi đã phát hiện và cách sửa: (1) Tránh nhầm chỉ số Julia 1-based trong DP: ô F(i, b) phải lưu ở table[i+1, b+1]; "
        "(2) Tra cứu chi phí phải thông qua Dict ánh xạ từ Order.id, không dùng vị trí đơn trong mảng; "
        "(3) Thao tác hoàn tác undo! phải chèn phần tử vào ĐẦU hàng đợi bằng enqueue_front! để bảo toàn đúng thứ tự ưu tiên."
    )
    doc.paragraphs[31].text = (
        "Một chức năng thay đổi ảnh hưởng đến chức năng sau như thế nào? BFS tính sai khoảng cách sẽ dẫn đến sai chi phí costs, "
        "làm cho DP chọn sai tập đơn hoặc vi phạm ngân sách. Sắp xếp merge_sort không ổn định sẽ làm xáo trộn thứ tự ưu tiên giao hàng "
        "và làm sai cấu trúc cây BST tra cứu. Thao tác giao/hoàn tác sai sẽ phá vỡ bất biến bảo toàn đơn và bảo toàn tổng thời gian/lợi ích."
    )

    # Q1 -> Q5
    doc.paragraphs[34].text = (
        "Trả lời: Dãy sắp giảm theo tỷ số b/τ làm xáo trộn thứ tự ID (không còn tăng dần). "
        "Tiền điều kiện của tìm kiếm nhị phân bắt buộc mảng phải được sắp xếp theo đúng khóa tìm kiếm (Order.id tăng dần) "
        "để tại mỗi bước so sánh có thể loại trừ chính xác nửa không chứa khóa cần tìm."
    )
    doc.paragraphs[36].text = (
        "Trả lời: Đơn 2 tới điểm 4: Tuyến đường BFS ngắn nhất là 1 -> 2 -> 4 gồm h = 2 cạnh. "
        "Thời gian đi và về là 2h = 2*2 = 4; thời gian phục vụ là s = 1 -> τ = 2h + s = 4 + 1 = 5. "
        "Nếu các cạnh mang thời gian khác nhau, BFS KHÔNG CÒN cho thời gian ngắn nhất vì BFS chỉ tối ưu số cạnh, "
        "khi đó phải sử dụng thuật toán Dijkstra."
    )
    doc.paragraphs[38].text = (
        "Trả lời: Ô F(2, 8) nằm tại table[2+1, 8+1] = table[3, 9] trong Julia. "
        "Hai phương án: Không lấy đơn 2 (F(1, 8) = table[2, 9] = 6) và Lấy đơn 2 (v_2 + F(1, 8 - 5) = 8 + table[2, 4] = 8 + 6 = 14). "
        "Lấy giá trị lớn nhất max(6, 14) = 14. Nhánh lấy phải dùng hàng trước (i-1) để đảm bảo mỗi đơn hàng chỉ được chọn tối đa một lần (tính chất 0/1)."
    )
    doc.paragraphs[40].text = (
        "Trả lời: Đơn 1 được rút khỏi đỉnh ngăn xếp completed và chèn ngược lại vào đầu hàng đợi pending bằng enqueue_front!; "
        "tổng thời gian spent giảm từ 3 về 0, tổng lợi ích earned giảm từ 6 về 0; hàng đợi khôi phục [1, 2], ngăn xếp rỗng []. "
        "Bất biến kiểm tra: Set(pending) ∩ Set(completed) = ∅ và Set(pending) ∪ Set(completed) = Set(selected_ids)."
    )
    doc.paragraphs[42].text = (
        "Trả lời: Độ phức tạp: merge_sort: O(n log n) thời gian, O(n) bộ nhớ; binary_search: O(log n) thời gian, O(1) bộ nhớ; "
        "BFS: O(V + E) thời gian, O(V) bộ nhớ; DP bảng đầy đủ: O(n * B) thời gian, O(n * B) bộ nhớ. "
        "Không thể khẳng định mọi BST đều tra cứu trong O(log n) vì chi phí tra cứu phụ thuộc chiều cao cây h (O(h)). "
        "Chiều cao O(log n) chỉ đạt được khi cây cân bằng. Nếu chèn khóa theo thứ tự tăng/giảm dần, cây suy biến thành danh sách liên kết "
        "với h = n, khi đó chi phí tra cứu trường hợp xấu nhất là O(n)."
    )

    doc.save('Phieu_thuc_hanh_CTDL_GT_2026_10_09_Da_dien.docx')
    doc.save('NguyenHoangNam_thuchanhbuoi05.docx')
    print("Done generating Phieu_thuc_hanh_CTDL_GT_2026_10_09_Da_dien.docx and NguyenHoangNam_thuchanhbuoi05.docx")

if __name__ == '__main__':
    fill_worksheet()
